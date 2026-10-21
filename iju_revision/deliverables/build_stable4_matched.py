import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "sTable4"

FONT = "Times New Roman"
title_font = Font(name=FONT, size=11, bold=True)
hdr_font = Font(name=FONT, size=11, bold=True)
body_font = Font(name=FONT, size=11, bold=False)
note_font = Font(name=FONT, size=10, bold=True)
abbr_font = Font(name=FONT, size=11, bold=False)

medium = Side(style="medium")
center = Alignment(horizontal="center", vertical="center")
left = Alignment(horizontal="left", vertical="center")

ws["A1"] = "Supplementary Table 4 (revised). Discriminatory ability of risk classification integrating IMDC risk and inflammatory markers — nested bootstrap (cutoff and model coefficients both re-derived within each of 1,000 resamples)"
ws["A1"].font = title_font
ws["A1"].alignment = left
ws.merge_cells("A1:E1")
ws.row_dimensions[1].height = 33.65

headers = ["Model", "Apparent C-index", "Optimism-corrected C-index", "95% CI", "Calibration slope"]
hdr_row = 3
for j, h in enumerate(headers, start=1):
    c = ws.cell(row=hdr_row, column=j, value=h)
    c.font = hdr_font
    c.alignment = center
    c.border = Border(top=medium)
ws.row_dimensions[hdr_row].height = 28

rows = [
    ("IMDC",           0.622, 0.621, "0.556–0.681", 1.090),
    ("IMDC-NLR",       0.673, 0.651, "0.584–0.715", 0.835),
    ("IMDC-dNLR",      0.689, 0.665, "0.596–0.731", 0.846),
    ("IMDC-MLR",       0.687, 0.665, "0.601–0.728", 0.871),
    ("IMDC-SIRI",      0.665, 0.641, "0.568–0.710", 0.859),
    ("IMDC-CRP",       0.698, 0.674, "0.606–0.732", 0.792),
    ("IMDC-NLR-MLR",   0.684, 0.664, "0.598–0.725", 0.793),
    ("IMDC-dNLR-MLR",  0.718, 0.698, "0.632–0.756", 0.850),
]

r = hdr_row + 1
last_row = hdr_row + len(rows)
for name, app, corr, ci, slope in rows:
    is_last = (r == last_row)
    bottom = Border(bottom=medium) if is_last else Border()
    vals = [name, app, corr, ci, slope]
    for j, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = body_font
        c.alignment = center
        c.border = bottom
        if j in (2, 3, 5):
            c.number_format = "0.000"
    r += 1

ws.column_dimensions["A"].width = 18.2
ws.column_dimensions["B"].width = 17.9
ws.column_dimensions["C"].width = 19.5
ws.column_dimensions["D"].width = 15.5
ws.column_dimensions["E"].width = 16.5

r += 1
ws.cell(row=r, column=1, value="Abbreviations").font = note_font
r += 1
ws.cell(row=r, column=1, value=(
    "RCC: renal cell carcinoma, IMDC: International Metastatic RCC Database Consortium,\n"
    "NLR: neutrophil-to-lymphocyte ratio, dNLR: derived neutrophil-to-lymphocyte ratio, \n"
    "MLR: monocyte-to-lymphocyte ratio, SIRI: systemic inflammatory response index, \n"
    "CRP: C-reactive protein"
)).font = abbr_font
ws.cell(row=r, column=1).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 59

r += 1
note = ws.cell(row=r, column=1, value=(
    "Note: 95% CI is the percentile interval of (apparent C − optimism_b) across 1,000 resamples in which both the "
    "ROC cutoff and the model coefficients were re-derived; this differs from the original table, which bootstrapped "
    "only the coefficients of an already-fixed cutoff. PLR, SII, and PIV are not shown (not independently associated "
    "with OS after IMDC adjustment in Supplementary Table 3)."
))
note.font = Font(name=FONT, size=9, italic=True, color="555555")
note.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
ws.row_dimensions[r].height = 45

wb.save("sTable4_nested_matched_format.xlsx")
print("saved")
